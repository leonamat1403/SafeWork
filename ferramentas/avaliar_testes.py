"""Avalia os testes de campo listados na planilha e preenche os resultados (Etapa 1.5).

Para cada linha da aba Testes (com `arquivo` preenchido), processa o vídeo – ou o trecho entre `inicio_s` e
`fim_s` – exatamente como o run.py (filtro do arco, modelos, regras e ocorrências) e grava:
  - na aba Testes: frames, FPS, % de cada estado, frames ignorados pelo arco, nº de ocorrências;
  - na aba Ocorrências: uma linha por ocorrência (a classificação VP/FP automática sai das fórmulas da planilha).

Uso (na pasta do projeto, com o ambiente .venv ativo e a planilha FECHADA no Excel):
    python ferramentas/avaliar_testes.py                  # avalia só as linhas ainda não avaliadas
    python ferramentas/avaliar_testes.py --todos          # reavalia todas as linhas
    python ferramentas/avaliar_testes.py --salvar-saida   # também grava os vídeos anotados em testes_campo/

Na primeira execução, se testes_campo/planilha_testes_campo.xlsx não existir, ela é criada a partir do modelo
docs/planilha_testes_campo.xlsx.

Importação automática: todo vídeo novo colocado em
    testes_campo/videos/conformes/      (pessoas COM máscara abaixada o tempo todo)  -> situacao_real = mascara_abaixada
    testes_campo/videos/nao_conformes/  (pessoas SEM máscara)                       -> situacao_real = sem_mascara
é cadastrado sozinho na aba Testes (e na aba Fontes) antes da avaliação. Revise depois o ângulo e as condições.
"""
from __future__ import annotations

import argparse
import os
import re
import shutil
import sys
import time
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))
os.environ.setdefault("YOLO_OFFLINE", "true")

import cv2  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from ferramentas.criar_planilha_testes import (COL, DESTINO_PADRAO, ULTIMA_LINHA,  # noqa: E402
                                               ULTIMA_LINHA_OCORRENCIAS, escrever_formulas_ocorrencia)
from src.config import Config, ErroConfig, carregar_config  # noqa: E402
from src.occurrences import GerenciadorOcorrencias, Ocorrencia, formatar_tempo  # noqa: E402
from src.tipos import Estado  # noqa: E402

PLANILHA_PADRAO = RAIZ / "testes_campo" / "planilha_testes_campo.xlsx"
PASTA_VIDEOS = RAIZ / "testes_campo" / "videos"
# Subpasta de testes_campo/videos -> situacao_real cadastrada automaticamente para os vídeos dela.
PASTAS_IMPORTACAO = {"conformes": "mascara_abaixada", "nao_conformes": "sem_mascara"}
EXTENSOES_VIDEO = {".mp4", ".avi", ".mov", ".mkv", ".webm", ".m4v"}


@dataclass
class ResultadoTeste:
    """Resumo da avaliação de um vídeo ou trecho: frames, FPS, fração de cada estado da máscara, frames ignorados
    pelo arco e as ocorrências geradas."""
    frames: int
    fps: float
    fracoes: dict[Estado, float]
    ignorados_arco: int
    ocorrencias: list[Ocorrencia] = field(default_factory=list)


def esperado(situacao: str | None, angulo: str | None) -> str:
    """Mesma regra da coluna 'esperado' da planilha."""
    if not situacao:
        return ""
    if situacao == "misto":
        return "revisão manual"
    if situacao in ("pausa_curta", "mascara_abaixada") or angulo == "costas":
        return "sem ocorrência"
    return "ocorrência"


def resultado(esperado_txt: str, ocorrencias: int | None) -> str:
    """Mesma regra da coluna 'resultado' da planilha."""
    if esperado_txt in ("", "revisão manual") or ocorrencias is None:
        return ""
    ok = (ocorrencias > 0 and esperado_txt == "ocorrência") or (ocorrencias == 0 and esperado_txt == "sem ocorrência")
    return "OK" if ok else "FALHA"


def avaliar_trecho(config: Config, caminho: Path, inicio_s: float | None, fim_s: float | None,
                   salvar_em: Path | None = None) -> ResultadoTeste:
    """Processa o vídeo (ou trecho) com o pipeline completo do sistema."""
    from src.detector import Detector
    from src.run import GravadorVideo, Pipeline, contar_estados, desenhar

    cap = cv2.VideoCapture(str(caminho))
    if not cap.isOpened():
        raise RuntimeError(f"não foi possível abrir {caminho}")
    # Começa no trecho pedido; o fim é controlado dentro do laço (tempo maior que fim_s).
    if inicio_s:
        cap.set(cv2.CAP_PROP_POS_MSEC, inicio_s * 1000)
    pipeline = Pipeline(config, Detector(config))       # detector novo por teste: o rastreamento começa do zero
    gerenciador = GerenciadorOcorrencias(config)
    gravador = GravadorVideo(salvar_em, cap.get(cv2.CAP_PROP_FPS), False) if salvar_em else None
    contagem: Counter = Counter()
    ocorrencias: list[Ocorrencia] = []
    frames = ignorados = 0
    inicio = time.perf_counter()
    try:
        while True:
            ok, frame = cap.read()
            if not ok:
                break
            t = cap.get(cv2.CAP_PROP_POS_MSEC) / 1000
            if fim_s is not None and t > fim_s:
                break
            r = pipeline.processar(frame)
            frames += 1
            ignorados += int(r.saturado)
            contar_estados(contagem, r)
            ocorrencias += gerenciador.atualizar(r.avaliacoes, t, frames, r.saturado)
            if gravador is not None:
                gravador.gravar(t, desenhar(frame, r, [f"{caminho.name} {formatar_tempo(t)}"]))
    finally:
        cap.release()
        if gravador is not None:
            gravador.fechar()
    duracao = time.perf_counter() - inicio
    # Fração de cada estado entre todas as avaliações de máscara (pessoa x frame).
    total = sum(n for (epi, _), n in contagem.items() if epi == "mascara")
    fracoes = {e: (contagem[("mascara", e)] / total if total else 0.0) for e in Estado}
    return ResultadoTeste(frames, frames / duracao if duracao > 0 else 0.0, fracoes, ignorados, ocorrencias)


def gravar_resultado(wb, n: int, id_teste: str, arquivo: str, res: ResultadoTeste) -> None:
    """Escreve o resultado na linha `n` da aba Testes e as ocorrências na aba Ocorrências."""
    ws = wb["Testes"]
    ws[f"{COL['frames']}{n}"] = res.frames
    ws[f"{COL['fps']}{n}"] = round(res.fps, 1)
    ws[f"{COL['pct_conforme']}{n}"] = round(res.fracoes[Estado.CONFORME], 3)
    ws[f"{COL['pct_nao_conforme']}{n}"] = round(res.fracoes[Estado.NAO_CONFORME], 3)
    ws[f"{COL['pct_indeterminado']}{n}"] = round(res.fracoes[Estado.INDETERMINADO], 3)
    ws[f"{COL['ignorados_arco']}{n}"] = res.ignorados_arco
    ws[f"{COL['ocorrencias']}{n}"] = len(res.ocorrencias)
    ws[f"{COL['avaliado_em']}{n}"] = datetime.now().strftime("%d/%m/%Y %H:%M")
    novas = [[id_teste, arquivo, formatar_tempo(oc.tempo_s), oc.track_id,
              "; ".join(f"{epi}: {motivo}" for epi, motivo in oc.situacao.items()), None] for oc in res.ocorrencias]
    substituir_ocorrencias(wb, id_teste, novas)


def substituir_ocorrencias(wb, id_teste: str, novas: list[list]) -> None:
    """Troca as linhas de um teste na aba Ocorrências, mantendo as dos outros testes (e a revisão manual delas).

    Se uma ocorrência nova for a mesma de antes (mesmo tempo no vídeo e mesma pessoa), a revisão manual dela é mantida;
    as diferentes ficam sem revisão, para serem conferidas de novo.
    """
    ws = wb["Ocorrências"]
    linhas, revisadas = [], {}
    for n in range(2, ULTIMA_LINHA_OCORRENCIAS + 1):
        # Colunas A-E (dados da ocorrência) e H (revisão manual); as demais (F, G, I) são fórmulas.
        valores = [ws.cell(n, col).value for col in (1, 2, 3, 4, 5, 8)]
        if valores[0] in (None, ""):
            continue
        if valores[0] != id_teste:
            linhas.append(valores)
        # Linhas do teste que será reavaliado: guarda a revisão manual para tentar reaproveitá-la.
        elif valores[5] not in (None, ""):
            revisadas[(valores[2], valores[3])] = valores[5]
    # Se a ocorrência nova coincide (mesmo tempo e pessoa) com uma antiga já revisada, herda a revisão.
    linhas += [linha[:5] + [revisadas.get((linha[2], linha[3]), linha[5])] for linha in novas]
    if len(linhas) > ULTIMA_LINHA_OCORRENCIAS - 1:
        raise RuntimeError("a aba Ocorrências está cheia; crie uma planilha nova para a próxima rodada de testes")
    for i in range(ULTIMA_LINHA_OCORRENCIAS - 1):
        n = i + 2
        valores = linhas[i] if i < len(linhas) else [None] * 6
        for col, valor in zip((1, 2, 3, 4, 5, 8), valores):
            ws.cell(n, col).value = valor
        if ws[f"F{n}"].value is None:     # garante as fórmulas (caso alguém tenha apagado a linha)
            escrever_formulas_ocorrencia(ws, n)


def importar_pastas(wb, pasta_videos: Path = PASTA_VIDEOS, raiz: Path = RAIZ) -> list[str]:
    """Cadastra na aba Testes (e em Fontes) os vídeos novos das pastas conformes/ e nao_conformes/."""
    ws = wb["Testes"]
    linhas_usadas = [n for n in range(2, ULTIMA_LINHA + 1)
                     if ws[f"{COL['arquivo']}{n}"].value or ws[f"{COL['ID']}{n}"].value]
    existentes = {str(ws[f"{COL['arquivo']}{n}"].value).replace("\\", "/") for n in linhas_usadas}
    # Números dos IDs já usados (T01, T02...); o próximo ID é o maior mais 1.
    numeros = [int(m.group(1)) for n in linhas_usadas
               if (m := re.fullmatch(r"T(\d+)", str(ws[f"{COL['ID']}{n}"].value or "")))]
    proximo = max(numeros, default=0) + 1
    # Linhas ainda vazias da aba Testes, em ordem.
    livres = (n for n in range(2, ULTIMA_LINHA + 1) if n not in set(linhas_usadas))
    fontes = wb["Fontes"]
    # Primeira linha vazia da aba Fontes (a aba tem até a linha 200).
    linha_fonte = next((n for n in range(2, 201) if not fontes[f"A{n}"].value), None)
    importados = []
    for subpasta, situacao in PASTAS_IMPORTACAO.items():
        pasta = pasta_videos / subpasta
        videos = sorted(p for p in pasta.glob("*") if p.suffix.lower() in EXTENSOES_VIDEO) if pasta.exists() else []
        for video in videos:
            relativo = video.relative_to(raiz).as_posix() if video.is_relative_to(raiz) else str(video)
            if relativo in existentes:
                continue
            n = next(livres, None)
            if n is None:
                raise RuntimeError("a aba Testes está cheia; crie uma planilha nova")
            id_teste = f"T{proximo:02d}"
            proximo += 1
            for coluna, valor in [("ID", id_teste), ("arquivo", relativo), ("situacao_real", situacao),
                                  ("angulo", "varia"), ("pessoas", None), ("origem", "internet"),
                                  ("observacoes", f"importado de videos/{subpasta} (revisar ângulo e condições)")]:
                ws[f"{COL[coluna]}{n}"] = valor
            if linha_fonte is not None and linha_fonte <= 200:
                # Colunas A (arquivo), B (origem) e F (uso) da aba Fontes.
                for col, valor in zip("ABF", (relativo, "internet", "teste")):
                    fontes[f"{col}{linha_fonte}"] = valor
                fontes[f"G{linha_fonte}"] = "preencher link e licença"
                linha_fonte += 1
            importados.append(id_teste)
    return importados


def main(argv: list[str] | None = None) -> int:
    """Avalia as linhas da aba Testes ainda sem resultado (todas, com --todos) e salva a planilha.
    Devolve 0 (ok), 2 (config inválida) ou 3 (planilha aberta no Excel)."""
    p = argparse.ArgumentParser(description="Avalia os testes de campo da planilha e preenche os resultados.")
    p.add_argument("--planilha", default=str(PLANILHA_PADRAO), help="planilha de testes (padrão: %(default)s)")
    p.add_argument("--config", default=str(RAIZ / "config.yaml"))
    p.add_argument("--posto", default="padrao")
    p.add_argument("--todos", action="store_true", help="reavalia também as linhas que já têm resultado")
    p.add_argument("--salvar-saida", action="store_true", help="grava os vídeos anotados em testes_campo/")
    args = p.parse_args(argv)

    planilha = Path(args.planilha)
    if not planilha.exists():
        planilha.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy(DESTINO_PADRAO, planilha)
        print(f"Planilha criada a partir do modelo: {planilha}")
    try:
        config = carregar_config(args.config, args.posto)
    except ErroConfig as erro:
        print(f"Configuração inválida: {erro}")
        return 2

    wb = load_workbook(planilha)
    importados = importar_pastas(wb)
    if importados:
        print(f"Vídeos novos cadastrados das pastas testes_campo/videos/: {', '.join(importados)}")
    ws = wb["Testes"]
    avaliados = 0
    for n in range(2, ULTIMA_LINHA + 1):
        id_teste = ws[f"{COL['ID']}{n}"].value
        arquivo = ws[f"{COL['arquivo']}{n}"].value
        if not arquivo:
            continue
        id_teste = str(id_teste or f"linha{n}")
        # Linha já avaliada (tem frames): pula, a menos que --todos.
        if ws[f"{COL['frames']}{n}"].value is not None and not args.todos:
            continue
        caminho = Path(arquivo) if Path(arquivo).is_absolute() else RAIZ / arquivo
        if not caminho.exists():
            print(f"[{id_teste}] arquivo não encontrado: {caminho} (linha ignorada)")
            continue
        inicio_s = ws[f"{COL['inicio_s']}{n}"].value
        fim_s = ws[f"{COL['fim_s']}{n}"].value
        situacao = ws[f"{COL['situacao_real']}{n}"].value
        angulo = ws[f"{COL['angulo']}{n}"].value
        if not situacao:
            print(f"[{id_teste}] ATENÇÃO: situacao_real vazia; o resultado (OK/FALHA) não será calculado.")
        trecho = f" ({inicio_s or 0}–{fim_s if fim_s is not None else 'fim'} s)" if inicio_s or fim_s else ""
        print(f"[{id_teste}] avaliando {arquivo}{trecho} ...", flush=True)
        saida = (RAIZ / "testes_campo" / f"avaliacao_{id_teste}_{caminho.stem}.mp4") if args.salvar_saida else None
        res = avaliar_trecho(config, caminho, float(inicio_s) if inicio_s else None,
                             float(fim_s) if fim_s is not None else None, saida)
        gravar_resultado(wb, n, id_teste, str(arquivo), res)
        esp = esperado(situacao, angulo)
        print(f"[{id_teste}] {res.frames} frames, {res.fps:.1f} FPS | CONFORME {res.fracoes[Estado.CONFORME]:.0%}, "
              f"NÃO CONFORME {res.fracoes[Estado.NAO_CONFORME]:.0%}, INDETERMINADO "
              f"{res.fracoes[Estado.INDETERMINADO]:.0%} | ocorrências: {len(res.ocorrencias)} | esperado: "
              f"{esp or '?'} -> {resultado(esp, len(res.ocorrencias)) or '-'}")
        avaliados += 1

    if avaliados == 0 and not importados:
        print("Nenhuma linha nova para avaliar. Preencha a aba Testes ou coloque vídeos em testes_campo/videos/"
              "conformes e nao_conformes (use --todos para reavaliar).")
        return 0
    wb.calculation.fullCalcOnLoad = True   # o Excel recalcula as métricas ao abrir
    try:
        wb.save(planilha)
    except PermissionError:
        print(f"\nNão foi possível salvar {planilha}: feche a planilha no Excel e rode de novo.")
        return 3
    print(f"\n{avaliados} teste(s) avaliado(s). Resultados salvos em {planilha} (abra a aba Métricas).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
