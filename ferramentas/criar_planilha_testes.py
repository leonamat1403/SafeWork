"""Cria o modelo da planilha de testes de campo: docs/planilha_testes_campo.xlsx (Etapa 1.5).

Abas:
  Instruções  – como preencher e como rodar a avaliação automática;
  Cenários    – o roteiro de testes (o que gravar), com o resultado esperado de cada cenário;
  Testes      – UMA LINHA POR VÍDEO (ou trecho de vídeo) testado: você preenche as colunas amarelas e o script
                ferramentas/avaliar_testes.py preenche as cinzas;
  Ocorrências – ocorrências registradas em cada teste, com coluna para a revisão manual (VP/FP);
  Fontes      – catálogo dos vídeos (origem, link, licença, autorização, uso);
  Métricas    – fórmulas que comparam os resultados com as METAS DE VALIDAÇÃO.

Uso:  python ferramentas/criar_planilha_testes.py  [destino.xlsx]
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = Path(__file__).resolve().parent.parent
DESTINO_PADRAO = RAIZ / "docs" / "planilha_testes_campo.xlsx"
ULTIMA_LINHA = 500            # linhas preparadas na aba Testes (fórmulas e listas)
ULTIMA_LINHA_OCORRENCIAS = 2000

FONTE = "Arial"
AMARELO = PatternFill("solid", fgColor="FFF2CC")      # você preenche
CINZA = PatternFill("solid", fgColor="EDEDED")        # o script preenche
AZUL_CABECALHO = PatternFill("solid", fgColor="1F4E78")
VERDE_CLARO = PatternFill("solid", fgColor="E2EFDA")
BORDA = Border(*(Side(style="thin", color="BFBFBF"),) * 4)

SITUACOES = ["sem_mascara", "mascara_levantada", "mascara_abaixada", "pausa_curta", "misto"]
ANGULOS = ["frente", "perfil", "costas", "varia"]
ARCO = ["apagado", "aceso"]
SIM_NAO = ["nao", "sim"]
ORIGENS = ["webcam", "celular", "internet", "laboratorio", "marcopolo"]

# Colunas da aba Testes: (cabeçalho, largura, preenchido por: "voce" | "script" | "formula", comentário)
COLUNAS_TESTES = [
    ("ID", 7, "voce", "Identificador do teste, ex.: T01."),
    ("arquivo", 42, "voce", "Caminho do vídeo a partir da pasta do projeto, ex.: testes_campo/c01_sem_mascara_2m.mp4"),
    ("inicio_s", 9, "voce", "Opcional: segundo do vídeo em que o trecho começa (vazio = início)."),
    ("fim_s", 8, "voce", "Opcional: segundo do vídeo em que o trecho termina (vazio = até o fim)."),
    ("cenario", 9, "voce", "Código do cenário da aba Cenários (ex.: C01)."),
    ("situacao_real", 18, "voce", "O que REALMENTE acontece no trecho: sem_mascara, mascara_levantada, mascara_abaixada, "
                                  "pausa_curta (tira/levanta a máscara por menos de 4 s) ou misto (vai para revisão manual)."),
    ("angulo", 9, "voce", "frente, perfil, costas ou varia."),
    ("distancia_m", 11, "voce", "Distância aproximada da câmera, em metros."),
    ("arco", 9, "voce", "apagado ou aceso."),
    ("fumos", 7, "voce", "nao ou sim."),
    ("pessoas", 8, "voce", "Quantas pessoas aparecem."),
    ("origem", 11, "voce", "webcam, celular, internet, laboratorio ou marcopolo."),
    ("observacoes", 28, "voce", "Livre."),
    ("esperado", 15, "formula", "Calculado a partir de situacao_real e angulo: 'ocorrência' ou 'sem ocorrência'."),
    ("frames", 8, "script", "Frames processados no trecho."),
    ("fps", 7, "script", "FPS do processamento completo (pose + máscara + regras) neste computador."),
    ("pct_conforme", 11, "script", "Fração dos frames (pessoa x frame) em CONFORME."),
    ("pct_nao_conforme", 13, "script", "Fração dos frames (pessoa x frame) em NÃO CONFORME."),
    ("pct_indeterminado", 14, "script", "Fração dos frames (pessoa x frame) em INDETERMINADO."),
    ("ignorados_arco", 12, "script", "Frames ignorados pelo filtro do arco."),
    ("ocorrencias", 11, "script", "Ocorrências registradas no trecho."),
    ("resultado", 10, "formula", "OK se o sistema fez o esperado (ocorrência quando devia, nenhuma quando não devia)."),
    ("avaliado_em", 17, "script", "Data e hora da avaliação."),
]
COL = {nome: get_column_letter(i + 1) for i, (nome, *_rest) in enumerate(COLUNAS_TESTES)}

CENARIOS = [
    # (ID, descrição, situação, ângulo, distância, arco, fumos, recurso necessário, prioridade)
    ("C01", "Sem máscara, de frente para a câmera", "sem_mascara", "frente", 2, "apagado", "nao", "webcam/celular", "alta"),
    ("C02", "Sem máscara, de perfil", "sem_mascara", "perfil", 2, "apagado", "nao", "webcam/celular", "alta"),
    ("C03", "Sem máscara, de costas (o sistema deve ficar AMARELO, sem ocorrência)", "sem_mascara", "costas", 2,
     "apagado", "nao", "webcam/celular", "média"),
    ("C04", "Sem máscara, de frente, mais longe", "sem_mascara", "frente", 4, "apagado", "nao", "celular", "alta"),
    ("C05", "Sem máscara, de frente, longe", "sem_mascara", "frente", 6, "apagado", "nao", "celular", "média"),
    ("C06", "Sem máscara, mexendo a cabeça (cima, baixo, lados, perfil)", "sem_mascara", "varia", 1, "apagado", "nao",
     "webcam", "alta"),
    ("C07", "Sem máscara com luz forte de lado ou contraluz", "sem_mascara", "perfil", 2, "apagado", "nao",
     "webcam/celular", "média"),
    ("C08", "Sem máscara de solda, mas de boné ou capacete de segurança", "sem_mascara", "frente", 2, "apagado", "nao",
     "webcam/celular", "média"),
    ("C09", "Pessoa passa sem máscara por menos de 4 s", "pausa_curta", "frente", 2, "apagado", "nao", "webcam/celular",
     "média"),
    ("C10", "Máscara abaixada, de frente", "mascara_abaixada", "frente", 2, "apagado", "nao", "máscara de solda", "alta"),
    ("C11", "Máscara abaixada, de perfil", "mascara_abaixada", "perfil", 2, "apagado", "nao", "máscara de solda", "alta"),
    ("C12", "Máscara abaixada, de costas", "mascara_abaixada", "costas", 2, "apagado", "nao", "máscara de solda", "média"),
    ("C13", "Máscara abaixada, a 4 m e a 6 m", "mascara_abaixada", "frente", 4, "apagado", "nao", "máscara de solda",
     "média"),
    ("C14", "Máscara levantada (sobre a cabeça) por mais de 4 s", "mascara_levantada", "frente", 2, "apagado", "nao",
     "máscara de solda", "alta"),
    ("C15", "Levanta a máscara por ~2 s para olhar o cordão e abaixa", "pausa_curta", "frente", 2, "apagado", "nao",
     "máscara de solda", "alta"),
    ("C16", "Soldando com máscara abaixada e ARCO ACESO", "mascara_abaixada", "frente", 2, "aceso", "nao",
     "laboratório, Marcopolo ou internet", "alta"),
    ("C17", "Soldando com máscara abaixada, com fumos e faíscas", "mascara_abaixada", "frente", 2, "aceso", "sim",
     "laboratório, Marcopolo ou internet", "média"),
    ("C18", "Arco aceso e soldador sem máscara: SOMENTE vídeos já existentes (internet). Nunca encenar!",
     "sem_mascara", "frente", 2, "aceso", "nao", "internet", "baixa"),
    ("C19", "Duas pessoas: uma com máscara e outra sem", "misto", "frente", 2, "apagado", "nao",
     "máscara de solda + 2 pessoas", "média"),
]

EXEMPLO_TESTE = {   # teste real feito em 12/10/2026 (webcam, sem máscara, cabeça em movimento)
    "ID": "T01", "arquivo": "testes_campo/teste_sem_mascara.mp4", "cenario": "C06", "situacao_real": "sem_mascara",
    "angulo": "varia", "distancia_m": 1, "arco": "apagado", "fumos": "nao", "pessoas": 1, "origem": "webcam",
    "observacoes": "Exemplo: teste real da webcam (modelo v3)", "frames": 1440, "fps": 14.9, "pct_conforme": 0.005,
    "pct_nao_conforme": 0.908, "pct_indeterminado": 0.087, "ignorados_arco": 0, "ocorrencias": 1,
    "avaliado_em": "13/10/2026 08:52",
}


def fonte(negrito: bool = False, cor: str = "000000", tamanho: int = 10) -> Font:
    return Font(name=FONTE, bold=negrito, color=cor, size=tamanho)


def cabecalho(ws, linha: int, titulos: list[str], larguras: list[int]) -> None:
    for i, (titulo, largura) in enumerate(zip(titulos, larguras), start=1):
        c = ws.cell(linha, i, titulo)
        c.font = fonte(True, "FFFFFF")
        c.fill = AZUL_CABECALHO
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDA
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = ws.cell(linha + 1, 1)


def lista(ws, valores: list[str], intervalo: str) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(valores) + '"', allow_blank=True)
    dv.error = "Escolha um valor da lista."
    ws.add_data_validation(dv)
    dv.add(intervalo)


def formula_esperado(n: int) -> str:
    sit, ang = f"{COL['situacao_real']}{n}", f"{COL['angulo']}{n}"
    return (f'=IF({sit}="","",IF({sit}="misto","revisão manual",'
            f'IF(OR({sit}="pausa_curta",{sit}="mascara_abaixada",{ang}="costas"),"sem ocorrência","ocorrência")))')


def formula_resultado(n: int) -> str:
    esp, oc = f"{COL['esperado']}{n}", f"{COL['ocorrencias']}{n}"
    return (f'=IF(OR({esp}="",{esp}="revisão manual",{oc}=""),"",'
            f'IF(OR(AND({oc}>0,{esp}="ocorrência"),AND({oc}=0,{esp}="sem ocorrência")),"OK","FALHA"))')


def aba_instrucoes(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Instruções"
    ws.column_dimensions["A"].width = 115
    linhas = [
        ("Planilha de testes de campo – Verificador de EPIs de Solda (Fase 1, máscara de solda)", True, 14),
        ("", False, 10),
        ("COMO USAR", True, 11),
        ("1. Grave os vídeos seguindo a aba Cenários (roteiro completo em docs/roteiro_testes_campo.md). "
         "Salve-os em testes_campo/.", False, 10),
        ("2. Na aba Testes, preencha uma linha por vídeo (ou por trecho de vídeo) nas colunas AMARELAS. "
         "A coluna situacao_real é a mais importante: é o que REALMENTE acontece no trecho.", False, 10),
        ("3. Feche esta planilha no Excel e rode:  python ferramentas/avaliar_testes.py", False, 10),
        ("   O script processa cada linha nova e preenche as colunas CINZAS e a aba Ocorrências.", False, 10),
        ("4. Abra a planilha de novo. Na aba Ocorrências, confira cada ocorrência no vídeo e, se a classificação "
         "automática estiver errada, escreva VP ou FP na coluna revisao_manual.", False, 10),
        ("5. A aba Métricas mostra o resultado contra as METAS DE VALIDAÇÃO. Ela vai para o relatório da Etapa 1.6.",
         False, 10),
        ("", False, 10),
        ("CORES", True, 11),
        ("Amarelo = você preenche.   Cinza = o script preenche.   Branco com fórmula = calculado (não editar).",
         False, 10),
        ("A linha T01 é um EXEMPLO real (vídeo da webcam de 12/10/2026). Pode mantê-la ou apagá-la.", False, 10),
        ("", False, 10),
        ("REGRAS DO SISTEMA (config.yaml)", True, 11),
        ("Ocorrência: sem máscara ou máscara levantada em pelo menos 75% dos frames de 4 segundos. "
         "Mesma pessoa: nova ocorrência só depois de 60 s.", False, 10),
        ("De costas ou com o rosto não visível, o sistema fica AMARELO (INDETERMINADO) e nunca registra ocorrência: "
         "por isso o esperado para 'costas' é 'sem ocorrência'.", False, 10),
        ("", False, 10),
        ("SEGURANÇA E PRIVACIDADE", True, 11),
        ("Nunca peça para alguém soldar sem máscara para gerar um teste: cenários com arco aceso e sem máscara só com "
         "vídeos já existentes (internet).", False, 10),
        ("Grave só pessoas que autorizaram. Os vídeos ficam em testes_campo/ (fora do git). No relatório, use números "
         "e imagens com rostos borrados.", False, 10),
    ]
    for i, (texto, negrito, tamanho) in enumerate(linhas, start=1):
        c = ws.cell(i, 1, texto)
        c.font = fonte(negrito, tamanho=tamanho)
        c.alignment = Alignment(wrap_text=True, vertical="top")


def aba_cenarios(wb: Workbook) -> None:
    ws = wb.create_sheet("Cenários")
    titulos = ["ID", "descricao", "situacao_real", "angulo", "distancia_m", "arco", "fumos", "esperado",
               "recurso_necessario", "prioridade", "testes_feitos"]
    cabecalho(ws, 1, titulos, [6, 62, 18, 9, 11, 9, 7, 15, 30, 10, 12])
    for n, cen in enumerate(CENARIOS, start=2):
        for col, valor in enumerate(cen[:7], start=1):
            ws.cell(n, col, valor)
        ws.cell(n, 8, formula_esperado(n).replace(f"{COL['situacao_real']}{n}", f"C{n}")
                .replace(f"{COL['angulo']}{n}", f"D{n}"))
        ws.cell(n, 9, cen[7])
        ws.cell(n, 10, cen[8])
        ws.cell(n, 11, f"=COUNTIF(Testes!${COL['cenario']}$2:${COL['cenario']}${ULTIMA_LINHA},A{n})")
        for col in range(1, 12):
            c = ws.cell(n, col)
            c.font = fonte(cor="C00000" if cen[0] == "C18" and col == 2 else "000000")
            c.border = BORDA
            c.alignment = Alignment(wrap_text=True, vertical="top")


def aba_testes(wb: Workbook) -> None:
    ws = wb.create_sheet("Testes")
    cabecalho(ws, 1, [c[0] for c in COLUNAS_TESTES], [c[1] for c in COLUNAS_TESTES])
    for i, (nome, _larg, quem, comentario) in enumerate(COLUNAS_TESTES, start=1):
        ws.cell(1, i).comment = Comment(f"{comentario}\n(Preenchido por: {quem})", "projeto")
    for n in range(2, ULTIMA_LINHA + 1):
        for i, (nome, _larg, quem, _com) in enumerate(COLUNAS_TESTES, start=1):
            c = ws.cell(n, i)
            c.font = fonte()
            c.border = BORDA
            if quem == "voce":
                c.fill = AMARELO
            elif quem == "script":
                c.fill = CINZA
        ws[f"{COL['esperado']}{n}"] = formula_esperado(n)
        ws[f"{COL['resultado']}{n}"] = formula_resultado(n)
        for nome in ("pct_conforme", "pct_nao_conforme", "pct_indeterminado"):
            ws[f"{COL[nome]}{n}"].number_format = "0.0%"
        ws[f"{COL['fps']}{n}"].number_format = "0.0"
    for nome, valor in EXEMPLO_TESTE.items():
        ws[f"{COL[nome]}2"] = valor
    lista(ws, SITUACOES, f"{COL['situacao_real']}2:{COL['situacao_real']}{ULTIMA_LINHA}")
    lista(ws, ANGULOS, f"{COL['angulo']}2:{COL['angulo']}{ULTIMA_LINHA}")
    lista(ws, ARCO, f"{COL['arco']}2:{COL['arco']}{ULTIMA_LINHA}")
    lista(ws, SIM_NAO, f"{COL['fumos']}2:{COL['fumos']}{ULTIMA_LINHA}")
    lista(ws, ORIGENS, f"{COL['origem']}2:{COL['origem']}{ULTIMA_LINHA}")


def aba_ocorrencias(wb: Workbook) -> None:
    ws = wb.create_sheet("Ocorrências")
    titulos = ["ID_teste", "arquivo", "tempo_video", "pessoa", "situacao_detectada", "esperado_no_teste",
               "classificacao_automatica", "revisao_manual", "classificacao_final"]
    cabecalho(ws, 1, titulos, [9, 42, 12, 8, 34, 16, 14, 14, 14])
    exemplo = ["T01", "testes_campo/teste_sem_mascara.mp4", "00:00:03.8", 1, "mascara: rosto visível sem máscara"]
    for n in range(2, ULTIMA_LINHA_OCORRENCIAS + 1):
        for col in range(1, 10):
            c = ws.cell(n, col)
            c.font = fonte()
            c.border = BORDA
            c.fill = AMARELO if col == 8 else (CINZA if col <= 5 else PatternFill())
        escrever_formulas_ocorrencia(ws, n)
    for col, valor in enumerate(exemplo, start=1):
        ws.cell(2, col, valor)
    lista(ws, ["VP", "FP"], f"H2:H{ULTIMA_LINHA_OCORRENCIAS}")


def escrever_formulas_ocorrencia(ws, n: int) -> None:
    """Fórmulas de uma linha da aba Ocorrências (usado também pelo script de avaliação)."""
    ws[f"F{n}"] = (f'=IF(A{n}="","",IFERROR(INDEX(Testes!${COL["esperado"]}$2:${COL["esperado"]}${ULTIMA_LINHA},'
                   f'MATCH(A{n},Testes!$A$2:$A${ULTIMA_LINHA},0)),""))')
    ws[f"G{n}"] = f'=IF(F{n}="ocorrência","VP",IF(F{n}="sem ocorrência","FP",""))'
    ws[f"I{n}"] = f'=IF(H{n}<>"",H{n},G{n})'


def aba_fontes(wb: Workbook) -> None:
    ws = wb.create_sheet("Fontes")
    titulos = ["arquivo", "origem", "link", "licenca", "autorizacao_das_pessoas", "uso", "observacoes"]
    cabecalho(ws, 1, titulos, [42, 12, 45, 22, 22, 10, 30])
    exemplo = ["testes_campo/teste_sem_mascara.mp4", "webcam", "", "próprio", "sim (o próprio autor)", "teste",
               "Exemplo"]
    for n in range(2, 201):
        for col in range(1, 8):
            c = ws.cell(n, col)
            c.font = fonte()
            c.border = BORDA
            c.fill = AMARELO
    for col, valor in enumerate(exemplo, start=1):
        ws.cell(2, col, valor)
    lista(ws, ORIGENS, "B2:B200")
    lista(ws, ["teste", "treino"], "F2:F200")


def aba_metricas(wb: Workbook) -> None:
    ws = wb.create_sheet("Métricas")
    t = lambda nome: f"Testes!${COL[nome]}$2:${COL[nome]}${ULTIMA_LINHA}"   # noqa: E731
    fin = f"'Ocorrências'!$I$2:$I${ULTIMA_LINHA_OCORRENCIAS}"
    ws.column_dimensions["A"].width = 62
    for col, larg in zip("BCDE", (14, 14, 18, 60)):
        ws.column_dimensions[col].width = larg
    ws["A1"] = "Métricas dos testes de campo – comparação com as METAS DE VALIDAÇÃO"
    ws["A1"].font = fonte(True, tamanho=13)
    cabecalho(ws, 3, ["Métrica", "Valor", "Meta", "Situação", "Como é calculada"], [62, 14, 14, 18, 60])
    ws.freeze_panes = None

    avaliados_ocorrencia = f'COUNTIFS({t("esperado")},"ocorrência",{t("ocorrencias")},"<>")'
    avaliados_sem = f'COUNTIFS({t("esperado")},"sem ocorrência",{t("ocorrencias")},"<>")'
    metricas = [
        ("Testes avaliados (com resultado OK ou FALHA)",
         f'=COUNTIF({t("resultado")},"OK")+COUNTIF({t("resultado")},"FALHA")', None, None,
         "Linhas da aba Testes já processadas pelo script (sem contar 'misto')."),
        ("Recall da ausência de máscara (testes em que devia haver ocorrência e houve)",
         f'=IFERROR(COUNTIFS({t("esperado")},"ocorrência",{t("resultado")},"OK")/{avaliados_ocorrencia},"")',
         0.70, ">=", "Testes 'ocorrência' com resultado OK ÷ testes 'ocorrência' avaliados."),
        ("Falsos positivos (ocorrências classificadas como FP)",
         f'=IFERROR(COUNTIF({fin},"FP")/(COUNTIF({fin},"FP")+COUNTIF({fin},"VP")),"")',
         0.30, "<=", "FP ÷ (VP + FP) na aba Ocorrências (classificação final: revisão manual ou automática)."),
        ("Testes sem ocorrência esperada que tiveram alarme falso",
         f'=IFERROR(COUNTIFS({t("esperado")},"sem ocorrência",{t("resultado")},"FALHA")/{avaliados_sem},"")',
         None, None, "Testes 'sem ocorrência' com FALHA ÷ testes 'sem ocorrência' avaliados (informativo)."),
        ("FPS mínimo (processamento completo, imgsz 640)",
         f'=IF(COUNT({t("fps")})=0,"",MIN({t("fps")}))', 5, ">=", "Menor FPS entre os testes (RNF01)."),
        ("FPS médio", f'=IF(COUNT({t("fps")})=0,"",AVERAGE({t("fps")}))', None, None, "Média dos testes."),
        ("% frames NÃO CONFORME nos testes em que devia haver ocorrência",
         f'=IFERROR(AVERAGEIFS({t("pct_nao_conforme")},{t("esperado")},"ocorrência",{t("ocorrencias")},"<>"),"")',
         None, None, "Média de pct_nao_conforme (quanto do tempo a falta foi reconhecida)."),
        ("% frames CONFORME (\"OK\" falso) nos testes em que devia haver ocorrência",
         f'=IFERROR(AVERAGEIFS({t("pct_conforme")},{t("esperado")},"ocorrência",{t("ocorrencias")},"<>"),"")',
         None, None, "Média de pct_conforme nos testes sem máscara / máscara levantada."),
        ("% frames CONFORME nos testes com máscara abaixada",
         f'=IFERROR(AVERAGEIFS({t("pct_conforme")},{t("situacao_real")},"mascara_abaixada",{t("ocorrencias")},"<>"),"")',
         None, None, "Quanto do tempo a máscara correta foi reconhecida."),
        ("Frames ignorados pelo arco (total)", f'=SUM({t("ignorados_arco")})', None, None, "Soma da coluna ignorados_arco."),
    ]
    for n, (nome, formula, meta, sentido, como) in enumerate(metricas, start=4):
        ws.cell(n, 1, nome)
        ws.cell(n, 2, formula)
        if meta is not None:
            ws.cell(n, 3, meta)
            comparacao = f"B{n}>=C{n}" if sentido == ">=" else f"B{n}<=C{n}"
            ws.cell(n, 4, f'=IF(B{n}="","sem dados",IF({comparacao},"atingida","abaixo da meta"))')
        ws.cell(n, 5, como)
        for col in range(1, 6):
            c = ws.cell(n, col)
            c.font = fonte()
            c.border = BORDA
            c.alignment = Alignment(wrap_text=True, vertical="top")
        formato = "0.0" if "FPS" in nome else ("0" if nome.startswith(("Testes avaliados", "Frames")) else "0.0%")
        ws.cell(n, 2).number_format = formato
        ws.cell(n, 3).number_format = formato
    ws.cell(n + 1, 1, "Metas: prompt do projeto (METAS DE VALIDAÇÃO, métricas práticas dos testes de campo).").font = \
        fonte(cor="595959")

    # Recortes por condição
    linha = n + 3
    ws.cell(linha, 1, "Por condição (recall = testes 'ocorrência' com OK ÷ testes 'ocorrência' avaliados)").font = \
        fonte(True, tamanho=11)
    linha += 1
    for i, titulo in enumerate(["Condição", "Testes avaliados", "Recall", "Alarmes falsos", ""], start=1):
        c = ws.cell(linha, i, titulo)
        c.font = fonte(True, "FFFFFF")
        c.fill = AZUL_CABECALHO
        c.border = BORDA
    dist = t("distancia_m")
    # (rótulo, critérios do COUNTIFS) – distância em faixas, para incluir valores como 1,5 m ou 3 m
    grupos = [("distância até 2 m", f'{dist},"<=2"'),
              ("distância de 2 a 4 m", f'{dist},">2",{dist},"<=4"'),
              ("distância acima de 4 m", f'{dist},">4"')] + \
             [(f"ângulo: {v}", f'{t("angulo")},"{v}"') for v in ANGULOS] + \
             [(f"arco {v}", f'{t("arco")},"{v}"') for v in ARCO]
    for rotulo, crit in grupos:
        linha += 1
        ws.cell(linha, 1, rotulo)
        ws.cell(linha, 2, f'=COUNTIFS({crit},{t("ocorrencias")},"<>")')
        ws.cell(linha, 3, f'=IFERROR(COUNTIFS({crit},{t("esperado")},"ocorrência",{t("resultado")},"OK")/'
                          f'COUNTIFS({crit},{t("esperado")},"ocorrência",{t("ocorrencias")},"<>"),"")')
        ws.cell(linha, 4, f'=COUNTIFS({crit},{t("esperado")},"sem ocorrência",{t("resultado")},"FALHA")')
        ws.cell(linha, 3).number_format = "0.0%"
        for col in range(1, 5):
            ws.cell(linha, col).font = fonte()
            ws.cell(linha, col).border = BORDA
            if linha % 2:
                ws.cell(linha, col).fill = VERDE_CLARO


def criar(destino: Path = DESTINO_PADRAO) -> Path:
    wb = Workbook()
    wb._named_styles["Normal"].font = Font(name=FONTE, size=10)   # fonte padrão das células vazias
    aba_instrucoes(wb)
    aba_cenarios(wb)
    aba_testes(wb)
    aba_ocorrencias(wb)
    aba_fontes(wb)
    aba_metricas(wb)
    wb.active = wb.sheetnames.index("Testes")
    wb.calculation.fullCalcOnLoad = True   # o Excel recalcula todas as fórmulas ao abrir
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


if __name__ == "__main__":
    caminho = criar(Path(sys.argv[1]) if len(sys.argv) > 1 else DESTINO_PADRAO)
    print(f"Planilha criada: {caminho}")
