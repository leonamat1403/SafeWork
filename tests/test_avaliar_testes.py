"""Planilha de testes de campo e script de avaliação (Etapa 1.5) – sem processar vídeos."""
from datetime import datetime

import pytest
from openpyxl import load_workbook

from ferramentas.avaliar_testes import ResultadoTeste, esperado, gravar_resultado, resultado
from ferramentas.criar_planilha_testes import COL, criar
from src.occurrences import Ocorrencia
from src.tipos import Estado


@pytest.mark.parametrize("situacao, angulo, esperado_txt", [
    ("sem_mascara", "frente", "ocorrência"),
    ("sem_mascara", "varia", "ocorrência"),
    ("mascara_levantada", "perfil", "ocorrência"),
    ("sem_mascara", "costas", "sem ocorrência"),         # de costas o sistema fica INDETERMINADO
    ("mascara_abaixada", "frente", "sem ocorrência"),
    ("pausa_curta", "frente", "sem ocorrência"),         # menos de 4 s em falta
    ("misto", "frente", "revisão manual"),
    (None, "frente", ""),
])
def test_resultado_esperado_de_cada_situacao(situacao, angulo, esperado_txt):
    assert esperado(situacao, angulo) == esperado_txt


def test_resultado_ok_ou_falha():
    assert resultado("ocorrência", 2) == "OK"
    assert resultado("ocorrência", 0) == "FALHA"
    assert resultado("sem ocorrência", 0) == "OK"
    assert resultado("sem ocorrência", 1) == "FALHA"
    assert resultado("revisão manual", 1) == ""
    assert resultado("ocorrência", None) == ""


def ocorrencia(t: float) -> Ocorrencia:
    return Ocorrencia(datetime.now(), t, 1, 3, ["mascara"], {"mascara": "rosto visível sem máscara"},
                      {"mascara": 1.0}, [], {}, False)


def resultado_simulado(n_ocorrencias: int) -> ResultadoTeste:
    return ResultadoTeste(300, 12.34, {Estado.CONFORME: 0.1, Estado.NAO_CONFORME: 0.8, Estado.INDETERMINADO: 0.1},
                          5, [ocorrencia(4.0 + 60 * i) for i in range(n_ocorrencias)])


def test_planilha_modelo(tmp_path):
    wb = load_workbook(criar(tmp_path / "modelo.xlsx"))
    assert wb.sheetnames == ["Instruções", "Cenários", "Testes", "Ocorrências", "Fontes", "Métricas"]
    assert wb["Testes"][f"{COL['esperado']}2"].value.startswith("=IF(")
    assert wb["Testes"][f"{COL['ID']}2"].value == "T01"          # linha de exemplo
    assert wb.calculation.fullCalcOnLoad


def test_importa_videos_das_pastas_conformes_e_nao_conformes(tmp_path):
    from ferramentas.avaliar_testes import importar_pastas
    for subpasta, nomes in {"conformes": ["a.mp4", "notas.txt"], "nao_conformes": ["b.MP4", "c.avi"]}.items():
        (tmp_path / "videos" / subpasta).mkdir(parents=True)
        for nome in nomes:
            (tmp_path / "videos" / subpasta / nome).write_bytes(b"x")
    wb = load_workbook(criar(tmp_path / "planilha.xlsx"))

    assert importar_pastas(wb, tmp_path / "videos", tmp_path) == ["T02", "T03", "T04"]
    ws = wb["Testes"]
    linhas = [(ws[f"{COL['ID']}{n}"].value, ws[f"{COL['arquivo']}{n}"].value, ws[f"{COL['situacao_real']}{n}"].value)
              for n in (3, 4, 5)]
    assert linhas == [("T02", "videos/conformes/a.mp4", "mascara_abaixada"),
                      ("T03", "videos/nao_conformes/b.MP4", "sem_mascara"),
                      ("T04", "videos/nao_conformes/c.avi", "sem_mascara")]
    assert wb["Fontes"]["A3"].value == "videos/conformes/a.mp4"
    assert importar_pastas(wb, tmp_path / "videos", tmp_path) == []        # não cadastra o mesmo vídeo duas vezes


def test_gravar_resultado_e_reavaliar_mantendo_a_revisao_dos_outros_testes(tmp_path):
    caminho = criar(tmp_path / "planilha.xlsx")
    wb = load_workbook(caminho)
    wb["Ocorrências"]["H2"] = "FP"                                  # revisão manual da ocorrência do exemplo T01
    gravar_resultado(wb, 3, "T02", "testes_campo/t02.mp4", resultado_simulado(2))
    ws = wb["Testes"]
    assert ws[f"{COL['frames']}3"].value == 300
    assert ws[f"{COL['fps']}3"].value == 12.3
    assert ws[f"{COL['pct_nao_conforme']}3"].value == 0.8
    assert ws[f"{COL['ignorados_arco']}3"].value == 5
    assert ws[f"{COL['ocorrencias']}3"].value == 2

    gravar_resultado(wb, 3, "T02", "testes_campo/t02.mp4", resultado_simulado(1))   # reavaliação do T02
    oc = wb["Ocorrências"]
    linhas = [[oc.cell(n, c).value for c in (1, 3, 4, 8)] for n in range(2, 10) if oc.cell(n, 1).value]
    assert linhas == [["T01", "00:00:03.8", 1, "FP"], ["T02", "00:00:04.0", 3, None]]
    assert oc["I3"].value.startswith("=IF(")                        # fórmulas da classificação continuam lá
